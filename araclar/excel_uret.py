"""EVALORA Excel dosyalarını CSV'lerden yeniden üretir.

Kullanım (EVALORA klasöründen):  python3 araclar/excel_uret.py
Gerekli paket: openpyxl  (python3 -m pip install openpyxl)

Her dersin CSV'leri kendi klasöründeki veri/ altındadır (ör. Fizik/veri); ortak kurallar ve
değişiklik günlüğü ortak/ altındadır. CSV'ler tek doğruluk kaynağıdır; Excel'i elle düzenleme,
CSV'yi düzenleyip betiği çalıştır. Betik tüm derslerin Excel'lerini, kökteki Ünite Kodları.xlsx'i
ve harita_uret.py ile mikro kazanım haritalarını, denetim raporlarını ve graf JSON'larını üretir.
"""

import csv
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import zorluk

KOK = Path(__file__).resolve().parent.parent
ORTAK = KOK / "ortak"

# Her ders: klasör, CSV öneki, Excel adı, ünite sayfası adı, ana ünite sayfası adı
DERSLER = [
    {"ad": "Fizik", "klasor": KOK / "Fizik", "onek": "fizik",
     "excel": "Fizik Mikro Kazanım.xlsx", "unite_sayfa": "Üniteler", "unite_csv": "fizik_uniteler.csv",
     "ana_sayfa": "Ana Üniteler", "ana_csv": "fizik_ana_uniteler.csv"},
    {"ad": "Matematik", "klasor": KOK / "Matematik", "onek": "matematik",
     "excel": "Matematik Mikro Kazanım.xlsx", "unite_sayfa": "Temalar", "unite_csv": "matematik_uniteler.csv",
     "ana_sayfa": "Ana Temalar", "ana_csv": "matematik_ana_temalar.csv"},
]

BASLIKLAR = {
    "kod": "Kod", "ad": "Ad", "siniflar": "Sınıflar", "sinif_uniteleri": "Sınıf Üniteleri",
    "sinif_temalari": "Sınıf Temaları",
    "sinif": "Sınıf", "unite_no": "Ünite No", "unite_adi": "Ünite Adı", "ana_unite": "Ana Ünite",
    "tymm_url": "TYMM Bağlantısı", "tema_no": "Tema No", "tema_adi": "Tema Adı",
    "tymm_unite_id": "TYMM Ünite No", "tymm_program_url": "TYMM Program Bağlantısı",
    "yanilgi": "Yanılgı", "ilgili_mikro_kazanimlar": "İlgili Mikro Kazanımlar",
    "ifade": "Mikro Kazanım (Öğrenci...)", "bilissel_duzey": "Bilişsel Düzey",
    "on_kosul_kodlari": "Ön Koşul (Mikro Kazanım)", "on_kosul_diger": "Ön Koşul (Diğer)",
    "yanilgilar": "Yanılgılar", "olcme_turleri": "Ölçme Türleri", "ustalik_olcutu": "Ustalık Ölçütü",
    "durum": "Durum", "eklenme_tarihi": "Eklenme Tarihi", "mikro_kod": "Mikro Kazanım Kodu",
    "sinif_unite": "Sınıf Ünitesi / Teması", "ogrenme_ciktisi": "Öğrenme Çıktısı (TYMM)",
    "surec_bileseni": "Süreç Bileşeni", "no": "No", "kural": "Kural", "ornek": "Örnek",
    "tarih": "Tarih", "degisiklik": "Değişiklik", "kisaltma": "Kısaltma",
    "ornek_mikro_kod": "Örnek Mikro Kod", "ana_unite_adi": "Ana Ünite Adı", "ana_tema": "Ana Tema",
    "tur": "Tür (Bilgi/Beceri)", "mikro_kazanim_sayisi": "Mikro Kazanım Sayısı",
    "islem_turu": "İşlem Türü", "islem_degeri": "İşlem Değeri", "alt_mk_sayisi": "Doğrudan Ön Koşul Sayısı",
    "zorluk": "Zorluk", "seviye": "Seviye", "deger": "Değer", "grup": "Grup", "tanim": "Tanım",
    "tipik_fiiller": "Tipik Fiiller", "ornekler": "Örnekler", "zorluk_alt": "Zorluk (alt sınır)",
    "zorluk_ust": "Zorluk (üst sınır)", "aciklama": "Açıklama", "sayisal_ornek": "Sayısal Ders Örneği",
    "sozel_ornek": "Sözel Ders Örneği",
}

GENIS = {"ifade", "yanilgi", "kural", "degisiklik", "ad", "unite_adi", "tema_adi",
         "on_kosul_diger", "ustalik_olcutu", "tymm_url", "tymm_program_url", "ornek",
         "tanim", "tipik_fiiller", "ornekler", "aciklama", "sayisal_ornek", "sozel_ornek"}

BASLIK_DOLGU = PatternFill("solid", fgColor="1F4E78")
BASLIK_YAZI = Font(bold=True, color="FFFFFF")


def csv_oku(yol):
    with open(yol, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def sayi(v):
    """Sayısal metni sayıya çevirir; "05" gibi baştaki sıfırlı kodlar metin olarak kalır."""
    if not isinstance(v, str) or not v or (v.startswith("0") and v != "0" and not v.startswith("0.")):
        return v
    try:
        return int(v)
    except ValueError:
        try:
            return float(v)
        except ValueError:
            return v


def zorluk_ekle(mikro):
    """Mikro kazanım satırlarına işlem değeri, doğrudan ön koşul sayısı, zorluk ve seviye sütunlarını ekler."""
    z = zorluk.hesapla(mikro)
    sonuc = []
    for m in mikro:
        r = dict(m)
        h = z[m["kod"]]
        r["islem_degeri"] = h["deger"] if h["deger"] is not None else "TANIMSIZ"
        r["alt_mk_sayisi"] = h["alt_sayisi"]
        r["zorluk"] = h["zorluk"]
        r["seviye"] = h["seviye"] if not h["eksik"] else "eksik işlem türü"
        sonuc.append(r)
    return sonuc


def sayfa_ekle(wb, baslik, satirlar):
    ws = wb.create_sheet(baslik)
    if not satirlar:
        ws.append(["(henüz kayıt yok)"])
        return
    alanlar = list(satirlar[0].keys())
    ws.append([BASLIKLAR.get(a, a) for a in alanlar])
    for s in satirlar:
        # "05" gibi baştaki sıfırlı kodlar metin olarak kalır
        ws.append([sayi(v) for v in (s[a] for a in alanlar)])
    for i, alan in enumerate(alanlar, start=1):
        harf = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[harf].width = 60 if alan in GENIS else 18
        for hucre in ws[harf]:
            hucre.alignment = Alignment(wrap_text=True, vertical="top")
    for hucre in ws[1]:
        hucre.fill = BASLIK_DOLGU
        hucre.font = BASLIK_YAZI
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def kitap_kaydet(yol, sayfalar):
    wb = Workbook()
    wb.remove(wb.active)
    for baslik, satirlar in sayfalar:
        sayfa_ekle(wb, baslik, satirlar)
    wb.save(yol)
    print(f"Üretildi: {yol.relative_to(KOK)}")


def mikro_sayilari(d):
    """Sınıf ünitesi/teması başına eşlenen (farklı) mikro kazanım sayısı."""
    yol = d["klasor"] / "veri" / f"{d['onek']}_kazanim_esleme.csv"
    if not yol.exists():
        return Counter()
    goruldu = {(r["sinif_unite"], r["mikro_kod"]) for r in csv_oku(yol)}
    return Counter(u for u, _ in goruldu)


if __name__ == "__main__":
    kurallar = csv_oku(ORTAK / "kodlama_kurallari.csv")
    gunluk = csv_oku(ORTAK / "degisiklik_gunlugu.csv")
    islem_turleri = csv_oku(ORTAK / "islem_turleri.csv")
    seviyeler = csv_oku(ORTAK / "zorluk_seviyeleri.csv")
    unite_sayfalari = []
    for d in DERSLER:
        veri = d["klasor"] / "veri"
        oku = lambda ad: csv_oku(veri / ad) if (veri / ad).exists() else []
        uniteler = oku(d["unite_csv"])
        sayilar = mikro_sayilari(d)
        for u in uniteler:
            u["mikro_kazanim_sayisi"] = sayilar.get(u["kod"], 0)
        unite_sayfalari.append((d["ad"], uniteler))
        if not (veri / f"{d['onek']}_mikro_kazanimlar.csv").exists():
            continue
        kitap_kaydet(d["klasor"] / d["excel"], [
            ("Mikro Kazanımlar", zorluk_ekle(oku(f"{d['onek']}_mikro_kazanimlar.csv"))),
            ("Kazanım Eşleme", oku(f"{d['onek']}_kazanim_esleme.csv")),
            ("Yanılgılar", oku(f"{d['onek']}_yanilgilar.csv")),
            (d["unite_sayfa"], uniteler),
            (d["ana_sayfa"], oku(d["ana_csv"])),
            ("İşlem Türleri", islem_turleri),
            ("Zorluk Seviyeleri", seviyeler),
            ("Kodlama Kuralları", kurallar),
            ("Değişiklik Günlüğü", gunluk),
        ])
    kitap_kaydet(KOK / "Ünite Kodları.xlsx", unite_sayfalari)

    # ilişki haritaları ve denetim raporları da güncellensin
    import harita_uret
    for d in harita_uret.DERSLER:
        harita_uret.uret(d)
